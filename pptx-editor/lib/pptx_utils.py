"""pptx-editor 共通ユーティリティ"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

# 1 cm = 360000 EMU
CM_TO_EMU = 360000


def cm(value: float) -> Emu:
    return Emu(int(value * CM_TO_EMU))


def emu_to_cm(emu: int) -> float:
    return emu / CM_TO_EMU


def load_prs(path: str) -> Presentation:
    return Presentation(path)


def save_prs(prs: Presentation, src_path: str, output: str | None) -> str:
    dest = output if output else src_path
    prs.save(dest)
    return dest


def parse_color(hex_str: str) -> RGBColor:
    """#RRGGBB または RRGGBB を RGBColor に変換"""
    h = hex_str.lstrip("#")
    r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    return RGBColor(r, g, b)


def get_slide(prs: Presentation, slide_num: int):
    """1-indexed のスライド番号でスライドを返す"""
    idx = slide_num - 1
    if idx < 0 or idx >= len(prs.slides):
        raise ValueError(f"スライド {slide_num} は存在しません（全{len(prs.slides)}枚）")
    return prs.slides[idx]


def find_shape(slide, name_or_index: str):
    """シェイプを名前（部分一致）またはインデックス（0始まり）で検索"""
    try:
        idx = int(name_or_index)
        return slide.shapes[idx]
    except ValueError:
        for shape in slide.shapes:
            if name_or_index in shape.name:
                return shape
        raise ValueError(f"シェイプ '{name_or_index}' が見つかりません")
