"""add-text コマンド — テキストボックスを追加する"""
import argparse
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from lib.pptx_utils import load_prs, save_prs, cm, get_slide, parse_color


def add_parser(subparsers):
    p = subparsers.add_parser("add-text", help="テキストボックスを追加")
    p.add_argument("file", help="pptxファイルパス")
    p.add_argument("--slide", type=int, required=True, help="スライド番号（1始まり）")
    p.add_argument("--text", required=True, help="テキスト内容")
    p.add_argument("--left", type=float, required=True, help="左端位置（cm）")
    p.add_argument("--top", type=float, required=True, help="上端位置（cm）")
    p.add_argument("--width", type=float, required=True, help="幅（cm）")
    p.add_argument("--height", type=float, required=True, help="高さ（cm）")
    p.add_argument("--font-size", type=float, default=None, help="フォントサイズ（pt）")
    p.add_argument("--font-name", default=None, help="フォント名")
    p.add_argument("--bold", action="store_true", help="太字")
    p.add_argument("--color", default=None, help="文字色 (#RRGGBB)")
    p.add_argument("--align", choices=["left", "center", "right"], default="left", help="文字揃え")
    p.add_argument("--output", default=None, help="出力先（省略時は上書き）")
    p.set_defaults(func=run)


_ALIGN_MAP = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def run(args: argparse.Namespace):
    prs = load_prs(args.file)
    slide = get_slide(prs, args.slide)

    txBox = slide.shapes.add_textbox(cm(args.left), cm(args.top), cm(args.width), cm(args.height))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.alignment = _ALIGN_MAP[args.align]

    run_ = p.add_run()
    run_.text = args.text

    font = run_.font
    if args.font_size is not None:
        font.size = Pt(args.font_size)
    if args.font_name:
        font.name = args.font_name
    if args.bold:
        font.bold = True
    if args.color:
        font.color.rgb = parse_color(args.color)

    dest = save_prs(prs, args.file, args.output)
    print(f"テキスト追加完了: {dest}  スライド{args.slide}")
