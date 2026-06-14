"""edit-table コマンド — 既存テーブルのセル内容・スタイルを編集する"""
import argparse
from pptx.util import Pt
from lib.pptx_utils import load_prs, save_prs, get_slide, find_shape, parse_color


def add_parser(subparsers):
    p = subparsers.add_parser("edit-table", help="テーブルのセルを編集")
    p.add_argument("file", help="pptxファイルパス")
    p.add_argument("--slide", type=int, required=True, help="スライド番号（1始まり）")
    p.add_argument("--shape", required=True, help="テーブルのシェイプ名（部分一致）または0始まりインデックス")
    p.add_argument("--row", type=int, required=True, help="行インデックス（0始まり）")
    p.add_argument("--col", type=int, required=True, help="列インデックス（0始まり）")
    p.add_argument("--text", default=None, help="セルに設定するテキスト")
    p.add_argument("--fill", default=None, help="背景色 (#RRGGBB)")
    p.add_argument("--font-size", type=float, default=None, help="フォントサイズ (pt)")
    p.add_argument("--bold", action="store_true", help="太字")
    p.add_argument("--color", default=None, help="文字色 (#RRGGBB)")
    p.add_argument("--output", default=None, help="出力先（省略時は上書き）")
    p.set_defaults(func=run)


def run(args: argparse.Namespace):
    prs = load_prs(args.file)
    slide = get_slide(prs, args.slide)
    shape = find_shape(slide, args.shape)

    if not shape.has_table:
        raise ValueError(f"シェイプ '{args.shape}' はテーブルではありません")

    table = shape.table
    rows = len(table.rows)
    cols = len(table.columns)
    if args.row < 0 or args.row >= rows:
        raise ValueError(f"行インデックス {args.row} が範囲外です（0〜{rows - 1}）")
    if args.col < 0 or args.col >= cols:
        raise ValueError(f"列インデックス {args.col} が範囲外です（0〜{cols - 1}）")

    cell = table.cell(args.row, args.col)

    if args.text is not None:
        cell.text = args.text

    if args.fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = parse_color(args.fill)

    if args.font_size is not None or args.bold or args.color:
        tf = cell.text_frame
        for para in tf.paragraphs:
            for run in para.runs:
                if args.font_size is not None:
                    run.font.size = Pt(args.font_size)
                if args.bold:
                    run.font.bold = True
                if args.color:
                    run.font.color.rgb = parse_color(args.color)

    dest = save_prs(prs, args.file, args.output)
    print(f"セル編集完了: {dest}  スライド{args.slide}  ({args.row},{args.col})")
